"""The document reader: types from the bytes, the ZIP limits, rows out of OCR boxes, each Office reader, a
PDF's text layer and layout mode, and, with the OCR engine installed, scanned and sideways pages read in
the child through the API. Every document here is made up in the test; no real person or company."""

import io
import zipfile
from collections.abc import AsyncIterator
from datetime import datetime
from pathlib import Path

import httpx
import pytest
from fastapi import FastAPI

from conftest import ocr_rec_dir
from niadra_models import documents, ocr
from niadra_models.api import create_app
from niadra_models.documents import DOCX, PDF, PPTX, XLSX, office_type, sniff
from niadra_models.settings import Settings

docx = pytest.importorskip("docx")
openpyxl = pytest.importorskip("openpyxl")
pptx = pytest.importorskip("pptx")
pytest.importorskip("pypdf")
pdfium = pytest.importorskip("pypdfium2")

OPTIONS = ocr.OcrOptions(1600, 960, 2, ocr_rec_dir())
NEEDS_OCR = pytest.mark.skipif(
    not ocr.available(OPTIONS.rec_dir), reason="the reader is the `ocr` extra and the pinned recognizer"
)


def saved(document: object) -> bytes:
    out = io.BytesIO()
    document.save(out)  # type: ignore[attr-defined]
    return out.getvalue()


def letter() -> bytes:
    d = docx.Document()
    d.add_paragraph("Northwind Home Appliances")
    d.add_paragraph("Case reference: CS-2026-78412")
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate(
        [("Item", "Qty", "Amount"), ("AquaPure 300 filter", "1", "$249.00"), ("", "", "")]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    merged = table.cell(2, 0).merge(table.cell(2, 1))
    merged.text = "Return shipping label"
    table.cell(2, 2).text = "$0.00"
    d.add_paragraph("Ship it with RMA-559013 on the box.")
    return saved(d)


def statement() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Statement"
    ws.append(["Account statement, Horizon Digital Bank"])
    ws.append([])
    ws.append(["Date", "Description", "Amount", None, None])
    ws.append([datetime(2026, 9, 8), "Power bill", -312.47])
    ws.append([datetime(2026, 9, 20, 14, 5), "Refund OS-55190", 152.0])
    ws.append(["Total", "=SUM(C4:C5)"])
    notes = wb.create_sheet("Notes")
    notes.append(["Closed on", "30/09/2026"])
    return saved(wb)


def slides() -> bytes:
    from pptx.util import Inches

    deck = pptx.Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Ticket CH-2026-30417"
    first.placeholders[1].text = "App failure reported on 16/09/2026\nFix shipped in 5.12.3"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Next steps"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    for r, (a, b) in enumerate([("Action", "Due"), ("Partial refund", "30/09/2026")]):
        table.cell(r, 0).text, table.cell(r, 1).text = a, b
    group = second.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(4), Inches(3), Inches(1)
    ).text_frame.text = "Visit ZIP 50030-230"
    return saved(deck)


def zipped(entries: dict[str, bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in entries.items():
            archive.writestr(name, data)
    return out.getvalue()


def rewritten(data: bytes, name: str, change: dict[bytes, bytes]) -> bytes:
    """The same ZIP with one entry's bytes changed."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = {i.filename: archive.read(i) for i in archive.infolist()}
    for old, new in change.items():
        entries[name] = entries[name].replace(old, new)
    return zipped(entries)


def test_each_type_is_told_by_its_bytes() -> None:
    assert sniff(b"%PDF-1.7\n") == PDF
    assert sniff(b"\xff\xd8\xff\xe0") == "image/jpeg"
    assert sniff(b"\x89PNG\r\n\x1a\n") == "image/png"
    assert (sniff(letter()), sniff(statement()), sniff(slides())) == (DOCX, XLSX, PPTX)
    assert sniff(b"GIF89a") is None
    assert sniff(zipped({"notes.txt": b"hello"})) is None  # a ZIP that is not an Office file
    assert sniff(b"PK\x03\x04 not really a zip") is None


def test_macro_enabled_files_are_not_read() -> None:
    # A .docm declares its main part macro-enabled.
    docm = rewritten(
        letter(),
        "[Content_Types].xml",
        {
            b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml": (
                b"application/vnd.ms-word.document.macroEnabled.main+xml"
            )
        },
    )
    assert sniff(docm) is None
    # A workbook that carries a VBA project under an .xlsx content type is refused all the same.
    with zipfile.ZipFile(io.BytesIO(statement())) as archive:
        entries = {i.filename: archive.read(i) for i in archive.infolist()}
    entries["xl/vbaProject.bin"] = b"\xd0\xcf\x11\xe0"
    assert sniff(zipped(entries)) is None


def test_a_zip_declaring_two_kinds_of_document_is_not_read() -> None:
    both = rewritten(
        letter(),
        "[Content_Types].xml",
        {
            b"</Types>": b'<Override PartName="/xl/workbook.xml" ContentType="application/vnd.'
            b'openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>'
        },
    )
    assert sniff(both) is None


def test_a_zip_bomb_is_refused_before_anything_inflates(monkeypatch: pytest.MonkeyPatch) -> None:
    data = letter()
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        total = sum(i.file_size for i in archive.infolist())
        count = len(archive.infolist())
    monkeypatch.setattr(documents, "ZIP_MAX_UNCOMPRESSED", total - 1)
    assert office_type(data) is None
    monkeypatch.setattr(documents, "ZIP_MAX_UNCOMPRESSED", total)
    assert office_type(data) == DOCX
    monkeypatch.setattr(documents, "ZIP_MAX_ENTRIES", count - 1)
    assert office_type(data) is None


def test_a_real_bomb_is_refused_by_its_declared_size() -> None:
    # 101 MiB of zeros deflate to about 100 KiB; the declared size gives it away.
    bomb = zipped({"[Content_Types].xml": b"<Types/>", "word/document.xml": b"\0" * (101 * 1024 * 1024)})
    assert len(bomb) < 1024 * 1024
    assert office_type(bomb) is None


def box(left: float, top: float, width: float, height: float) -> list[list[float]]:
    return [[left, top], [left + width, top], [left + width, top + height], [left, top + height]]


def test_boxes_on_one_visual_line_become_one_row_left_to_right() -> None:
    result = [
        (box(900, 402, 120, 30), "R$ 612,00", 0.9),
        (box(100, 400, 400, 30), "Compressor EMX 1/4 HP", 0.9),
        (box(700, 405, 20, 30), "1", 0.9),
        (box(100, 300, 500, 32), "ORDEM DE SERVICO OS-55190", 0.9),
        (box(100, 460, 300, 30), "Termostato digital", 0.9),
        (box(900, 470, 120, 30), "R$ 89,90", 0.9),  # a third of a line lower still joins its row
    ]
    assert ocr.rows(result) == [
        "ORDEM DE SERVICO OS-55190",
        "Compressor EMX 1/4 HP  1  R$ 612,00",
        "Termostato digital  R$ 89,90",
    ]
    assert ocr.rows(None) == []


def test_a_box_more_than_half_a_line_away_starts_a_row() -> None:
    result = [(box(0, 0, 100, 20), "top", 0.9), (box(200, 11, 100, 20), "next", 0.9)]
    assert ocr.rows(result) == ["top", "next"]


def test_docx_keeps_the_body_order_and_one_line_per_table_row() -> None:
    found = documents.read(letter(), DOCX, OPTIONS)
    assert found.source == "office"
    assert found.text.splitlines() == [
        "Northwind Home Appliances",
        "Case reference: CS-2026-78412",
        "Item | Qty | Amount",
        "AquaPure 300 filter | 1 | $249.00",
        "Return shipping label | $0.00",  # the merged cell once
        "Ship it with RMA-559013 on the box.",
    ]


def test_xlsx_reads_cached_values_of_every_sheet_and_never_a_formula() -> None:
    found = documents.read(statement(), XLSX, OPTIONS)
    assert (found.source, found.pages) == ("office", 2)
    assert found.text.splitlines() == [
        "Statement",
        "Account statement, Horizon Digital Bank",
        "Date | Description | Amount",
        "2026-09-08 | Power bill | -312.47",
        "2026-09-20 14:05:00 | Refund OS-55190 | 152",
        "Total",  # a formula openpyxl never computed has no cached value, and its text stays out
        "Notes",
        "Closed on | 30/09/2026",
    ]


def test_xlsx_is_opened_read_only_with_cached_values(monkeypatch: pytest.MonkeyPatch) -> None:
    import openpyxl as module

    seen: dict[str, object] = {}
    original = module.load_workbook

    def spy(*args: object, **kwargs: object) -> object:
        seen.update(kwargs)
        return original(*args, **kwargs)  # type: ignore[arg-type]

    monkeypatch.setattr(module, "load_workbook", spy)
    documents.read(statement(), XLSX, OPTIONS)
    assert (seen["read_only"], seen["data_only"], seen["keep_links"]) == (True, True, False)


def test_pptx_reads_slides_in_order_with_tables_and_groups() -> None:
    found = documents.read(slides(), PPTX, OPTIONS)
    assert (found.source, found.pages) == ("office", 2)
    assert found.text.splitlines() == [
        "Ticket CH-2026-30417",
        "App failure reported on 16/09/2026",
        "Fix shipped in 5.12.3",
        "Next steps",
        "Action | Due",
        "Partial refund | 30/09/2026",
        "Visit ZIP 50030-230",
    ]


def test_a_long_document_is_read_whole() -> None:
    # Past 100,000 characters the rest of a letter was left out. Every line is read.
    d = docx.Document()
    for i in range(5000):
        d.add_paragraph(f"Line {i} of a long letter")
    found = documents.read(saved(d), DOCX, OPTIONS)
    lines = found.text.splitlines()
    assert len(lines) == 5000
    assert lines[-1] == "Line 4999 of a long letter"
    assert found.next_page is None


def test_a_damaged_office_file_is_unreadable_not_a_crash() -> None:
    damaged = rewritten(letter(), "word/document.xml", {b"<w:body>": b"<w:body"})
    with pytest.raises(documents.UnreadableDocumentError):
        documents.read(damaged, DOCX, OPTIONS)


def native_pdf(
    texts: list[tuple[float, float, str]], rects: list[tuple[float, float, float, float]]
) -> bytes:
    """A one-page PDF with a text layer and, optionally, ruled cells, written by hand."""
    ops = [f"{x} {y} {w} {h} re S" for x, y, w, h in rects]
    ops += [f"BT /F1 11 Tf {x} {y} Td ({t}) Tj ET" for x, y, t in texts]
    stream = "\n".join(ops).encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = b"%PDF-1.4\n", []
    for i, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{off:010d} 00000 n \n".encode() for off in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return out


INVOICE_ROWS = [("Router AX3000", "1", "R$ 489,90"), ("Cable Cat6 5 m", "3", "R$ 73,50")]


def invoice(ruled: bool) -> bytes:
    texts = [(72, 720, "INVOICE 2026-004817")]
    rects = []
    for r, row in enumerate(INVOICE_ROWS):
        y = 680 - 24 * r
        for x, w, cell in zip((72, 272, 332), (200, 60, 120), row, strict=True):
            # Written cell by cell, column after column: the text layer alone splits the rows.
            texts.append((x + 4, y + 7, cell))
            rects.append((x, y, w, 24))
    texts.sort(key=lambda t: (t[0], -t[1]))
    return native_pdf(texts, rects if ruled else [])


def test_a_ruled_table_keeps_each_row_on_one_line() -> None:
    found = documents.read(invoice(ruled=True), PDF, OPTIONS)
    assert (found.source, found.pages, found.ocr_pages) == ("pdf_text", 1, 0)
    lines = [" ".join(line.split()) for line in found.text.splitlines()]
    assert "Router AX3000 1 R$ 489,90" in lines
    assert "Cable Cat6 5 m 3 R$ 73,50" in lines


def test_a_page_without_rules_keeps_the_plain_text_layer() -> None:
    plain = documents.read(invoice(ruled=False), PDF, OPTIONS).text
    assert "INVOICE 2026-004817" in plain
    assert "Router AX3000 1 R$ 489,90" not in [" ".join(line.split()) for line in plain.splitlines()]


def test_a_huge_page_is_rendered_within_the_pixel_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    doc = pdfium.PdfDocument.new()
    doc.new_page(14400, 14400)  # 200 inches square: 1.6 billion pixels at 200 dpi
    monkeypatch.setattr(ocr, "MAX_PIXELS", 4_000_000)
    image = documents._render(doc, 0)
    assert image.width * image.height <= 4_000_000
    assert image.width > 1900
    small = pdfium.PdfDocument.new()
    small.new_page(612, 792)
    assert documents._render(small, 0).size == (1700, 2200)  # a letter page at 200 dpi


def test_an_encrypted_or_unopenable_pdf_gives_no_text_or_unreadable() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter(clone_from=io.BytesIO(invoice(ruled=False)))
    writer.encrypt("secret", algorithm="RC4-128")
    locked = io.BytesIO()
    writer.write(locked)
    assert documents.read(locked.getvalue(), PDF, OPTIONS).text == ""
    with pytest.raises(documents.UnreadableDocumentError):
        documents.read(b"%PDF-1.7 not a pdf at all", PDF, OPTIONS)


SERVICE_ORDER = [
    ("SERVICE ORDER OS-55190", "", ""),
    ("Opened on 03/09/2026", "", ""),
    ("Protocol 2026090312345", "", ""),
    ("Part or service", "Qty", "Price"),
    ("Compressor EMX 1/4 HP", "1", "R$ 612,00"),
    ("Digital thermostat", "1", "R$ 89,90"),
    ("Technician labour", "1", "R$ 180,00"),
    ("Quote total: R$ 881,90", "", ""),
]


def scanned_page(rotate: int = 0) -> object:
    from PIL import Image, ImageDraw, ImageFont

    page = Image.new("L", (1654, 2339), 250)
    draw = ImageDraw.Draw(page)
    font = ImageFont.load_default(size=34)
    for i, row in enumerate(SERVICE_ORDER):
        for x, cell in zip((120, 900, 1150), row, strict=True):
            draw.text((x, 150 + 80 * i), cell, font=font, fill=25)
    return page.rotate(rotate, expand=True) if rotate else page


def image_pdf(*pages: object) -> bytes:
    out = io.BytesIO()
    first, *rest = pages
    first.save(out, "PDF", resolution=200, save_all=True, append_images=rest)  # type: ignore[attr-defined]
    return out.getvalue()


def mixed_pdf() -> bytes:
    """Page 1 with a text layer, page 2 a scan lying on its side."""
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(native_pdf([(72, 720, "RETURN FORM DEV-2026-11873")], []))))
    writer.append(PdfReader(io.BytesIO(image_pdf(scanned_page(rotate=90)))))
    out = io.BytesIO()
    writer.write(out)
    return out.getvalue()


def flat(text: str) -> str:
    return "".join(text.split())


@pytest.fixture
async def client() -> AsyncIterator[httpx.AsyncClient]:
    app: FastAPI = create_app(Settings(ocr_rec_dir=OPTIONS.rec_dir, onnx_threads=2))
    async with (
        app.router.lifespan_context(app),
        httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://models") as c,
    ):
        yield c


async def read_through_the_server(client: httpx.AsyncClient, data: bytes) -> dict[str, object]:
    mime = sniff(data)
    assert mime is not None
    r = await client.post("/v1/document-text", params={"type": mime}, content=data, timeout=120)
    assert r.status_code == 200, r.text
    body: dict[str, object] = r.json()
    return body


@NEEDS_OCR
async def test_a_scanned_pdf_is_read_by_ocr_with_its_table_rows_on_one_line(
    client: httpx.AsyncClient,
) -> None:
    body = await read_through_the_server(client, image_pdf(scanned_page()))
    assert (body["source"], body["pages"], body["ocr_pages"]) == ("pdf_ocr", 1, 1)
    text = str(body["text"])
    for key in ("OS-55190", "03/09/2026", "2026090312345", "R$881,90"):
        assert key in flat(text), key
    rows = [flat(line) for line in text.splitlines()]
    assert any("CompressorEMX1/4HP" in r and "R$612,00" in r for r in rows)
    assert any("Digitalthermostat" in r and "R$89,90" in r for r in rows)


@NEEDS_OCR
async def test_a_page_scanned_sideways_is_read_after_the_text_layer(client: httpx.AsyncClient) -> None:
    body = await read_through_the_server(client, mixed_pdf())
    assert (body["source"], body["pages"], body["ocr_pages"]) == ("pdf_ocr", 2, 1)
    text = flat(str(body["text"]))
    assert text.startswith("RETURNFORMDEV-2026-11873")
    for key in ("OS-55190", "03/09/2026", "2026090312345", "R$612,00", "R$881,90"):
        assert key in text, key


@NEEDS_OCR
async def test_office_files_are_read_through_the_same_child(client: httpx.AsyncClient) -> None:
    letter_text = (await read_through_the_server(client, letter()))["text"]
    assert "AquaPure 300 filter | 1 | $249.00" in str(letter_text).splitlines()
    sheet_text = (await read_through_the_server(client, statement()))["text"]
    assert "2026-09-08 | Power bill | -312.47" in str(sheet_text).splitlines()


def test_the_image_build_selfcheck_reads_one_file_of_each_type() -> None:
    if not ocr.available(OPTIONS.rec_dir):
        pytest.skip("the reader is the `ocr` extra and the pinned recognizer")
    found = documents.selfcheck(OPTIONS)
    assert [line.split(": ")[1] for line in found] == ["ocr", "pdf_ocr", "office", "office", "office"]


def test_rec_dir_default_matches_the_image(tmp_path: Path) -> None:
    assert Settings().ocr_rec_dir == Path("/models/pp-ocrv6-small-rec-r1")
    assert not ocr.available(tmp_path)


def test_a_document_past_one_window_is_read_whole_a_window_at_a_time(monkeypatch: pytest.MonkeyPatch) -> None:
    # Past 50 sheets or slides the rest of a document was left out. One request reads a
    # window and says where the next starts, so the cell reads every one.
    whole = documents.read(statement(), XLSX, OPTIONS)
    assert whole.next_page is None
    monkeypatch.setattr(documents, "MAX_PAGES", 1)
    for document, mime in ((statement(), XLSX), (slides(), PPTX)):
        texts, first = [], 0
        while True:
            window = documents.read(document, mime, OPTIONS, first)
            texts.append(window.text)
            if window.next_page is None:
                break
            assert window.next_page == first + 1
            first = window.next_page
        assert len(texts) > 1
        monkeypatch.setattr(documents, "MAX_PAGES", 50)
        assert "\n".join(texts) == documents.read(document, mime, OPTIONS).text
        monkeypatch.setattr(documents, "MAX_PAGES", 1)


def test_a_pdf_is_read_a_window_of_pages_at_a_time(monkeypatch: pytest.MonkeyPatch) -> None:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for n in range(3):
        writer.append(PdfReader(io.BytesIO(native_pdf([(72, 720, f"PAGE NUMBER {n} OF THREE")], []))))
    out = io.BytesIO()
    writer.write(out)
    monkeypatch.setattr(documents, "MAX_PAGES", 2)
    first = documents.read(out.getvalue(), PDF, OPTIONS)
    assert (first.pages, first.next_page) == (2, 2)
    rest = documents.read(out.getvalue(), PDF, OPTIONS, 2)
    assert (rest.pages, rest.next_page) == (1, None)
    assert "PAGE NUMBER 2" in rest.text
    assert "PAGE NUMBER 2" not in first.text
