"""Synthetic customer-support documents (PT and EN) with their ground truth: the test set of the document
reader's eval. No real people or companies. Seeded, so a rerun writes the same set.

    uv run --with reportlab --with numpy python evals/documents/generate.py [--stress]

Needs the Arial and Courier New TrueType files (macOS has them; elsewhere set NIADRA_EVAL_FONTS to a
directory holding Arial.ttf, "Arial Bold.ttf" and "Courier New.ttf"). Writes testset/ next to this file,
which git ignores. `--stress` adds the two long documents of the memory measurement: 36 scanned pages
(about 19 MB, near the 20 MB limit) and 50 native pages.
"""

import io
import json
import os
import random
import sys
from pathlib import Path

import numpy as np
import pypdfium2 as pdfium
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from pptx import Presentation
from pptx.util import Inches
from pptx.util import Pt as PPt
from reportlab.lib.pagesizes import A4, letter
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

OUT = Path(__file__).parent / "testset"
OUT.mkdir(exist_ok=True)
FONTS = os.environ.get("NIADRA_EVAL_FONTS", "/System/Library/Fonts/Supplemental").rstrip("/") + "/"
pdfmetrics.registerFont(TTFont("Arial", FONTS + "Arial.ttf"))
pdfmetrics.registerFont(TTFont("Arial-Bold", FONTS + "Arial Bold.ttf"))
pdfmetrics.registerFont(TTFont("Courier", FONTS + "Courier New.ttf"))
random.seed(7)
np.random.seed(7)
TRUTH: dict[str, dict] = {}


def table(c, x, y, widths, rows, size=9.5, row_h=7 * mm):
    for r, row in enumerate(rows):
        c.setFont("Arial-Bold" if r == 0 else "Arial", size)
        cx = x
        for w, cell in zip(widths, row, strict=True):
            c.rect(cx, y - row_h, w, row_h)
            c.drawString(cx + 2 * mm, y - row_h + 2.2 * mm, cell)
            cx += w
        y -= row_h
    return y


# 1. Native-text invoice with a table (PT)
INV_ROWS = [
    ["Item", "Descrição", "Qtd", "Valor unit.", "Total"],
    ["1", "Roteador Wi-Fi 6 AX3000", "1", "R$ 489,90", "R$ 489,90"],
    ["2", "Cabo de rede Cat6 5 m", "3", "R$ 24,50", "R$ 73,50"],
    ["3", "Instalação técnica residencial", "1", "R$ 150,00", "R$ 150,00"],
    ["4", "Plano Fibra 600 Mega (agosto)", "1", "R$ 129,90", "R$ 129,90"],
]


def invoice(c):
    _, h = A4
    c.setFont("Arial-Bold", 16)
    c.drawString(20 * mm, h - 25 * mm, "Lojas Aurora Comércio Ltda")
    c.setFont("Arial", 9.5)
    c.drawString(
        20 * mm, h - 31 * mm, "Rua das Acácias, 1250, Funcionários, Belo Horizonte MG, CEP 30140-071"
    )
    c.setFont("Arial-Bold", 13)
    c.drawString(20 * mm, h - 45 * mm, "NOTA FISCAL DE SERVIÇO Nº 2026-004817")
    c.setFont("Arial", 10)
    c.drawString(20 * mm, h - 53 * mm, "Emitida em 14/08/2026      Vencimento 29/08/2026")
    c.drawString(
        20 * mm, h - 59 * mm, "Cliente: Marina Teixeira Duarte      Protocolo de atendimento 88213-554"
    )
    y = table(c, 20 * mm, h - 70 * mm, [12 * mm, 78 * mm, 14 * mm, 30 * mm, 30 * mm], INV_ROWS)
    c.setFont("Arial-Bold", 11)
    c.drawString(126 * mm, y - 8 * mm, "Total a pagar: R$ 843,30")
    c.setFont("Arial", 9.5)
    c.drawString(20 * mm, y - 22 * mm, "Pagamento via boleto bancário até a data de vencimento.")
    c.drawString(20 * mm, y - 28 * mm, "Dúvidas: central de atendimento 0800 555 0199, de segunda a sábado.")


c = canvas.Canvas(str(OUT / "invoice_pt.pdf"), pagesize=A4)
invoice(c)
c.save()
TRUTH["invoice_pt.pdf"] = {
    "kind": "native PDF with a table",
    "lang": "pt",
    "pages": 1,
    "keys": [
        "2026-004817",
        "14/08/2026",
        "29/08/2026",
        "30140-071",
        "88213-554",
        "R$ 843,30",
        "Roteador Wi-Fi 6 AX3000",
        "Cabo de rede Cat6",
        "Instalação técnica residencial",
        "Plano Fibra 600 Mega",
        "R$ 489,90",
        "R$ 73,50",
        "0800 555 0199",
    ],
    "order": [
        "Lojas Aurora",
        "NOTA FISCAL",
        "Roteador",
        "Plano Fibra",
        "Total a pagar",
        "Pagamento via boleto",
    ],
    "table_rows": [r[1:] for r in INV_ROWS[1:]],
}


# Helpers to fake a scan
def render_pdf_page(pdf_bytes: bytes, index: int, dpi: int) -> Image.Image:
    doc = pdfium.PdfDocument(pdf_bytes)
    img = doc[index].render(scale=dpi / 72).to_pil().convert("L")
    doc.close()
    return img


def scanify(img: Image.Image, angle: float, noise: float = 10, blur: float = 0.6) -> Image.Image:
    img = img.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=245)
    arr = np.asarray(img).astype(np.float32)
    arr = arr * 0.92 + 12 + np.random.normal(0, noise, arr.shape)
    img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(blur))
    buf = io.BytesIO()
    img.save(buf, "JPEG", quality=72)
    return Image.open(io.BytesIO(buf.getvalue()))


def image_pdf(pages: list[Image.Image], path: Path, pagesize=A4):
    c = canvas.Canvas(str(path), pagesize=pagesize)
    for img in pages:
        buf = io.BytesIO()
        img.save(buf, "JPEG", quality=75)
        buf.seek(0)
        from reportlab.lib.utils import ImageReader

        w, h = pagesize
        iw, ih = img.size
        s = min(w / iw, h / ih)
        c.drawImage(ImageReader(buf), (w - iw * s) / 2, (h - ih * s) / 2, iw * s, ih * s)
        c.showPage()
    c.save()


# 2. Scanned service order (PT)
OS_ROWS = [
    ["Peça ou serviço", "Qtd", "Valor"],
    ["Compressor EMX 1/4 HP", "1", "R$ 612,00"],
    ["Termostato digital", "1", "R$ 89,90"],
    ["Mão de obra técnica", "1", "R$ 180,00"],
]
buf = io.BytesIO()
c = canvas.Canvas(buf, pagesize=A4)
w, h = A4
c.setFont("Arial-Bold", 15)
c.drawString(20 * mm, h - 25 * mm, "ORDEM DE SERVIÇO OS-55190")
c.setFont("Arial", 10.5)
lines = [
    "Assistência Técnica Polar Refrigeração",
    "Av. Paulista, 1578, Bela Vista, São Paulo SP, CEP 01310-930",
    "Data de abertura: 03/09/2026      Protocolo: 2026090312345",
    "Equipamento: Refrigerador Frost Free 375 L, série PF375-88120",
    "Defeito relatado: não gela no compartimento inferior.",
]
for i, t in enumerate(lines):
    c.drawString(20 * mm, h - (36 + 7 * i) * mm, t)
y = table(c, 20 * mm, h - 80 * mm, [90 * mm, 20 * mm, 40 * mm], OS_ROWS, size=10.5, row_h=8 * mm)
c.setFont("Arial-Bold", 12)
c.drawString(20 * mm, y - 10 * mm, "Total do orçamento: R$ 881,90")
c.setFont("Arial", 10.5)
c.drawString(20 * mm, y - 20 * mm, "Garantia de 90 dias para peças e serviço.")
c.save()
scan = scanify(render_pdf_page(buf.getvalue(), 0, 200), angle=0.8)
image_pdf([scan], OUT / "service_order_scan_pt.pdf")
TRUTH["service_order_scan_pt.pdf"] = {
    "kind": "scanned PDF (image only)",
    "lang": "pt",
    "pages": 1,
    "keys": [
        "OS-55190",
        "01310-930",
        "03/09/2026",
        "2026090312345",
        "Refrigerador Frost Free 375 L",
        "PF375-88120",
        "Compressor EMX 1/4 HP",
        "Termostato digital",
        "R$ 612,00",
        "R$ 89,90",
        "R$ 180,00",
        "R$ 881,90",
    ],
    "order": [
        "ORDEM DE SERVIÇO",
        "Data de abertura",
        "Defeito relatado",
        "Compressor",
        "Total do orçamento",
        "Garantia",
    ],
    "table_rows": [r for r in OS_ROWS[1:]],
}


# 3. Photo-like receipt (EN), JPEG
RECEIPT = [
    ("GREENLEAF MARKET", None),
    ("Store #0412", None),
    ("123 Harbor Street", None),
    ("Portland, OR 97209", None),
    ("09/12/2026  18:42", None),
    ("", None),
    ("ORGANIC BANANAS", "$1.74"),
    ("ALMOND MILK 64OZ", "$4.29"),
    ("SOURDOUGH LOAF", "$5.49"),
    ("FREE RANGE EGGS 12CT", "$6.99"),
    ("", None),
    ("SUBTOTAL", "$18.51"),
    ("TAX", "$1.48"),
    ("TOTAL", "$19.99"),
    ("", None),
    ("TRANS 4471-0092-3318", None),
    ("THANK YOU FOR SHOPPING", None),
]
font = ImageFont.truetype(FONTS + "Courier New.ttf", 34)
paper = Image.new("L", (760, 60 + 52 * len(RECEIPT)), 250)
d = ImageDraw.Draw(paper)
for i, (left, right) in enumerate(RECEIPT):
    y = 30 + 52 * i
    d.text((40, y), left, font=font, fill=25)
    if right:
        d.text((720 - d.textlength(right, font=font), y), right, font=font, fill=25)
bg = Image.new("L", (1200, 1500), 120)
arr = np.asarray(bg).astype(np.float32) + np.linspace(-30, 30, 1500)[:, None]
bg = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
rot = paper.rotate(-2.5, resample=Image.BICUBIC, expand=True, fillcolor=0)
mask = Image.new("L", paper.size, 255).rotate(-2.5, expand=True, fillcolor=0)
bg.paste(rot, (200, 60), mask)
# mild perspective: squeeze the top
coeffs = (1.0, 0.03, -30, 0.0, 1.0, 0, 0.00001, 0.0)
photo = bg.transform(bg.size, Image.PERSPECTIVE, coeffs, Image.BICUBIC, fillcolor=110)
arr = np.asarray(photo).astype(np.float32) + np.random.normal(0, 8, (1500, 1200))
photo = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8)).filter(ImageFilter.GaussianBlur(1.0))
photo = Image.merge("RGB", [photo, photo.point(lambda v: v * 0.97), photo.point(lambda v: v * 0.9)])
photo.save(OUT / "receipt_photo_en.jpg", "JPEG", quality=70)
TRUTH["receipt_photo_en.jpg"] = {
    "kind": "receipt photo (JPEG)",
    "lang": "en",
    "pages": 1,
    "keys": [
        "GREENLEAF MARKET",
        "#0412",
        "97209",
        "09/12/2026",
        "ORGANIC BANANAS",
        "ALMOND MILK 64OZ",
        "SOURDOUGH LOAF",
        "FREE RANGE EGGS 12CT",
        "$18.51",
        "$1.48",
        "$19.99",
        "4471-0092-3318",
    ],
    "order": ["GREENLEAF", "Portland", "ORGANIC BANANAS", "SOURDOUGH", "SUBTOTAL", "TOTAL", "THANK YOU"],
    "table_rows": [
        ["ORGANIC BANANAS", "$1.74"],
        ["ALMOND MILK 64OZ", "$4.29"],
        ["SOURDOUGH LOAF", "$5.49"],
        ["FREE RANGE EGGS 12CT", "$6.99"],
    ],
}


# 4. DOCX letter (EN)
doc = Document()
doc.styles["Normal"].font.name = "Arial"
doc.styles["Normal"].font.size = Pt(11)
doc.add_paragraph("Northwind Home Appliances\n4410 Commerce Blvd, Austin, TX 78744")
doc.add_paragraph("September 18, 2026")
doc.add_paragraph("Case reference: CS-2026-78412")
doc.add_paragraph("Dear Mr. Daniel Okafor,")
doc.add_paragraph(
    "Thank you for contacting our support team about your AquaPure 300 water filter. We confirm that the "
    "unit is covered by the two-year warranty and that a full refund has been approved."
)
t = doc.add_table(rows=3, cols=2)
t.style = "Table Grid"
for r, (a, b) in enumerate(
    [("Item", "Amount"), ("AquaPure 300 water filter", "$249.00"), ("Return shipping label", "$0.00")]
):
    t.cell(r, 0).text, t.cell(r, 1).text = a, b
doc.add_paragraph(
    "Please ship the unit with return authorization RMA-559013 printed on the box. The refund of $249.00 "
    "will reach your card within 7 business days after we receive it."
)
doc.add_paragraph("Sincerely,\nCustomer Care Team")
doc.save(OUT / "letter_en.docx")
TRUTH["letter_en.docx"] = {
    "kind": "DOCX letter",
    "lang": "en",
    "pages": 1,
    "keys": [
        "78744",
        "September 18, 2026",
        "CS-2026-78412",
        "AquaPure 300 water filter",
        "$249.00",
        "RMA-559013",
        "7 business days",
    ],
    "order": [
        "Northwind",
        "September 18",
        "Case reference",
        "Dear",
        "two-year warranty",
        "RMA-559013",
        "Sincerely",
    ],
    "table_rows": [["AquaPure 300 water filter", "$249.00"], ["Return shipping label", "$0.00"]],
}


# 5. XLSX statement (PT)
wb = Workbook()
ws = wb.active
ws.title = "Extrato"
ST_ROWS = [
    ["Data", "Descrição", "Valor (R$)", "Saldo (R$)"],
    ["05/09/2026", "PIX recebido Loja Aurora", "1.250,00", "5.097,50"],
    ["08/09/2026", "Boleto Energia Clara", "-312,47", "4.785,03"],
    ["12/09/2026", "Compra cartão Mercado Sol", "-89,90", "4.695,13"],
    ["15/09/2026", "Tarifa pacote de serviços", "-29,90", "4.665,23"],
    ["20/09/2026", "PIX recebido Reembolso OS-55190", "152,50", "4.817,73"],
]
ws.append(["Extrato de conta, Banco Horizonte Digital"])
ws.append(["Av. Paulista, 1000, São Paulo SP, CEP 01310-100"])
ws.append(["Período: 01/09/2026 a 30/09/2026"])
ws.append([])
for row in ST_ROWS:
    ws.append(row)
ws.append([])
ws.append(["Saldo final", "", "", "R$ 4.817,73"])
wb.save(OUT / "statement_pt.xlsx")
TRUTH["statement_pt.xlsx"] = {
    "kind": "XLSX statement",
    "lang": "pt",
    "pages": 1,
    "keys": [
        "01310-100",
        "01/09/2026",
        "30/09/2026",
        "PIX recebido Loja Aurora",
        "Boleto Energia Clara",
        "-312,47",
        "Mercado Sol",
        "4.817,73",
        "OS-55190",
        "1.250,00",
    ],
    "order": ["Extrato de conta", "Período", "Loja Aurora", "Energia Clara", "Reembolso", "Saldo final"],
    "table_rows": [r for r in ST_ROWS[1:]],
}


# 6. Two-column PDF (EN), text drawn row by row across both columns
LEFT = [
    "Section A. Your plan changes on the",
    "effective date: October 1, 2026. The",
    "new plan code is PRM-240.",
    "",
    "Section B. The monthly price becomes",
    "$34.99 per month, taxes included, and",
    "the first bill shows a prorated amount.",
    "",
    "Section C. Devices bought before the",
    "change keep their current warranty.",
]
RIGHT = [
    "Section D. To cancel without a fee,",
    "open a request before September 30",
    "and keep the confirmation number.",
    "",
    "Section E. Written notices go to our",
    "office at 200 Park Avenue, New York,",
    "NY, ZIP 10017.",
    "",
    "Section F. Questions about this notice",
    "use support ticket TKT-883021.",
]
c = canvas.Canvas(str(OUT / "two_column_en.pdf"), pagesize=letter)
w, h = letter
c.setFont("Arial-Bold", 15)
c.drawString(20 * mm, h - 22 * mm, "Service Terms Update for Subscribers")
c.setFont("Arial", 10.5)
for i in range(max(len(LEFT), len(RIGHT))):
    y = h - (36 + 6 * i) * mm
    if i < len(LEFT) and LEFT[i]:
        c.drawString(20 * mm, y, LEFT[i])
    if i < len(RIGHT) and RIGHT[i]:
        c.drawString(112 * mm, y, RIGHT[i])
c.setFont("Arial", 8.5)
c.drawString(20 * mm, 18 * mm, "Brightwave Telecom, customer notice 2026/09")
c.save()
TRUTH["two_column_en.pdf"] = {
    "kind": "two-column native PDF",
    "lang": "en",
    "pages": 1,
    "keys": ["October 1, 2026", "PRM-240", "$34.99", "10017", "TKT-883021", "September 30"],
    "order": ["Section A", "Section B", "Section C", "Section D", "Section E", "Section F"],
    "contiguous": [
        "Your plan changes on the effective date: October 1, 2026",
        "The monthly price becomes $34.99 per month",
        "Written notices go to our office at 200 Park Avenue",
        "Questions about this notice use support ticket TKT-883021",
    ],
    "table_rows": [],
}


# 7. PDF with a native page and a rotated scanned page (PT)
buf = io.BytesIO()
c = canvas.Canvas(buf, pagesize=A4)
c.setFont("Arial-Bold", 15)
c.drawString(20 * mm, h - 25 * mm, "COMPROVANTE DE COLETA")
c.setFont("Arial", 11)
for i, t in enumerate(
    [
        "Transportadora Rota Sul Logística",
        "Rua XV de Novembro, 700, Centro, Curitiba PR, CEP 80020-000",
        "Código de rastreio: BR482915736BR",
        "Data da coleta: 22/09/2026",
        "Produto: Notebook Lumen 14 i5 16GB",
        "Valor declarado: R$ 1.349,00",
        "Volumes: 1      Peso: 2,4 kg",
    ]
):
    c.drawString(20 * mm, h - (38 + 8 * i) * mm, t)
c.save()
scanned = scanify(render_pdf_page(buf.getvalue(), 0, 200), angle=90.4)  # page scanned sideways
c = canvas.Canvas(str(OUT / "_p1.pdf"), pagesize=A4)
c.setFont("Arial-Bold", 15)
c.drawString(20 * mm, h - 25 * mm, "TERMO DE DEVOLUÇÃO DEV-2026-11873")
c.setFont("Arial", 11)
c.drawString(20 * mm, h - 38 * mm, "Solicitação aberta em 21/09/2026 pelo canal WhatsApp.")
c.drawString(
    20 * mm, h - 46 * mm, "Motivo: produto com defeito na tela. O comprovante de coleta segue anexo."
)
c.save()
image_pdf([scanned], OUT / "_p2.pdf", pagesize=(A4[1], A4[0]))
merged = pdfium.PdfDocument.new()
for part in ("_p1.pdf", "_p2.pdf"):
    src = pdfium.PdfDocument(str(OUT / part))
    merged.import_pages(src)
    src.close()
merged.save(str(OUT / "rotated_scan_pt.pdf"))
merged.close()
(OUT / "_p1.pdf").unlink()
(OUT / "_p2.pdf").unlink()
TRUTH["rotated_scan_pt.pdf"] = {
    "kind": "PDF with a page scanned sideways (90 degrees)",
    "lang": "pt",
    "pages": 2,
    "keys": [
        "DEV-2026-11873",
        "21/09/2026",
        "80020-000",
        "BR482915736BR",
        "22/09/2026",
        "Notebook Lumen 14 i5 16GB",
        "R$ 1.349,00",
        "Rota Sul",
    ],
    "order": [
        "TERMO DE DEVOLUÇÃO",
        "Motivo",
        "COMPROVANTE DE COLETA",
        "Código de rastreio",
        "Valor declarado",
    ],
    "table_rows": [],
}


# 8. PPTX (PT)
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Resumo do chamado CH-2026-30417"
s.placeholders[1].text = (
    "Cliente relatou falha no aplicativo em 16/09/2026\n"
    "Plano Empresarial Plus, R$ 1.980,00 por mês\n"
    "Correção publicada na versão 5.12.3"
)
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = "Próximos passos"
shape = s2.shapes.add_table(3, 2, Inches(1), Inches(2), Inches(7), Inches(1.5)).table
for r, (a, b) in enumerate(
    [
        ("Ação", "Prazo"),
        ("Reembolso proporcional", "30/09/2026"),
        ("Visita técnica em Recife PE, CEP 50030-230", "02/10/2026"),
    ]
):
    shape.cell(r, 0).text, shape.cell(r, 1).text = a, b
    for cell in (shape.cell(r, 0), shape.cell(r, 1)):
        cell.text_frame.paragraphs[0].runs[0].font.size = PPt(16)
prs.save(OUT / "summary_pt.pptx")
TRUTH["summary_pt.pptx"] = {
    "kind": "PPTX slides",
    "lang": "pt",
    "pages": 2,
    "keys": [
        "CH-2026-30417",
        "16/09/2026",
        "R$ 1.980,00",
        "5.12.3",
        "30/09/2026",
        "50030-230",
        "02/10/2026",
        "Reembolso proporcional",
    ],
    "order": [
        "Resumo do chamado",
        "Plano Empresarial",
        "Próximos passos",
        "Reembolso proporcional",
        "Visita técnica",
    ],
    "table_rows": [
        ["Reembolso proporcional", "30/09/2026"],
        ["Visita técnica em Recife PE, CEP 50030-230", "02/10/2026"],
    ],
}


# Long documents for the memory measurement (--stress); not scored
def stress() -> None:
    pages = []
    for n in range(36):
        out = io.BytesIO()
        c = canvas.Canvas(out, pagesize=A4)
        c.setFont("Arial-Bold", 15)
        c.drawString(20 * mm, h - 25 * mm, f"ORDEM DE SERVIÇO OS-{56000 + n}")
        c.setFont("Arial", 10.5)
        for i, t in enumerate(lines):
            c.drawString(20 * mm, h - (36 + 7 * i) * mm, t)
        y = table(c, 20 * mm, h - 80 * mm, [90 * mm, 20 * mm, 40 * mm], OS_ROWS, size=10.5, row_h=8 * mm)
        c.setFont("Arial-Bold", 12)
        c.drawString(20 * mm, y - 10 * mm, "Total do orçamento: R$ 881,90")
        c.save()
        pages.append(scanify(render_pdf_page(out.getvalue(), 0, 200), angle=0.2 * (n % 5)))
    image_pdf(pages, OUT / "stress_scan_36_pages_pt.pdf")
    c = canvas.Canvas(str(OUT / "stress_native_50_pages_pt.pdf"), pagesize=A4)
    for _ in range(50):
        invoice(c)
        c.showPage()
    c.save()


if "--stress" in sys.argv[1:]:
    stress()

(OUT / "ground_truth.json").write_text(json.dumps(TRUTH, ensure_ascii=False, indent=2), encoding="utf-8")
print("\n".join(f"{p.name} {p.stat().st_size}" for p in sorted(OUT.iterdir())))
